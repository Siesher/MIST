import shutil
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))
import assemble_vkr as av


@pytest.mark.skipif(
    shutil.which("pandoc") is None,
    reason="pandoc не установлен в системе; тест проверяет корректность локатора при наличии бинаря",
)
def test_find_pandoc_returns_existing_exe():
    p = av.find_pandoc()
    assert p, "pandoc не найден"
    assert Path(p).exists()
    assert "pandoc" in Path(p).name.lower()


def test_merge_fragment_replaces_placeholder_with_content():
    from docx import Document

    skel = Document()
    skel.add_paragraph("3.5. GSPO", style="Heading 2")
    skel.add_paragraph("[placeholder: что писать]", style="Normal")
    skel.add_paragraph("3.6. KTO", style="Heading 2")
    frag = Document()
    frag.add_paragraph("Первый абзац контента.", style="Normal")
    frag.add_paragraph("Второй абзац.", style="Normal")

    av.merge_fragment(skel, "3.5. GSPO", frag)
    texts = [p.text for p in skel.paragraphs]
    assert "[placeholder: что писать]" not in texts
    i35 = texts.index("3.5. GSPO")
    assert texts[i35 + 1] == "Первый абзац контента."
    assert texts[i35 + 2] == "Второй абзац."
    assert texts[i35 + 3] == "3.6. KTO"

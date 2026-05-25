"""Сборка ВКР: skeleton (python-docx) + главы (pandoc md->docx с OMML) -> merge.

Usage: uv run python scripts/assemble_vkr.py
"""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).parent))
from build_vkr_skeleton import OUT, SECTION_IDS, TEMPLATE, build_document  # noqa: E402

ROOT = Path(__file__).parent.parent
CHAPTERS = ROOT / "docs/diploma/chapters"


def find_pandoc() -> str:
    """Локатор pandoc: PATH -> %LOCALAPPDATA%\\Pandoc -> Program Files."""
    found = shutil.which("pandoc")
    if found:
        return found
    candidates = [
        Path(os.environ.get("LOCALAPPDATA", "")) / "Pandoc" / "pandoc.exe",
        Path(os.environ.get("ProgramFiles", "")) / "Pandoc" / "pandoc.exe",
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    raise RuntimeError("pandoc не найден (PATH/LOCALAPPDATA/ProgramFiles)")


def merge_fragment(doc, heading_prefix: str, fragment_doc) -> int:
    """Вставляет body фрагмента после заголовка heading_prefix, удаляя placeholder.

    Возвращает число вставленных элементов. Raises если заголовок не найден.
    """
    body = doc.element.body
    heading_el = None
    for p in doc.paragraphs:
        if p.text.strip().startswith(heading_prefix):
            heading_el = p._element
            break
    if heading_el is None:
        raise ValueError(f"заголовок {heading_prefix!r} не найден")
    # удалить placeholder-параграф сразу после заголовка (текст начинается с '[')
    nxt = heading_el.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        if Paragraph(nxt, doc).text.strip().startswith("["):
            body.remove(nxt)
    # вставить элементы фрагмента (кроме sectPr) после заголовка
    anchor = heading_el
    inserted = 0
    for child in list(fragment_doc.element.body):
        if child.tag == qn("w:sectPr"):
            continue
        new = deepcopy(child)
        anchor.addnext(new)
        anchor = new
        inserted += 1
    return inserted


def md_to_fragment(md_path: Path, pandoc: str) -> Path:
    """pandoc md -> временный docx с OMML-формулами и стилями образца."""
    frag = Path(tempfile.gettempdir()) / f"_vkr_frag_{md_path.stem}.docx"
    subprocess.run(
        [
            pandoc,
            str(md_path),
            "--shift-heading-level-by=1",
            f"--reference-doc={TEMPLATE}",
            "-o",
            str(frag),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return frag


# Сопоставление маркеров [Рис ...] с доступными изображениями (substring -> [(путь, подпись)])
FIGURES = [
    (
        "скриншоты интерфейса",
        [
            (
                "figures/fig_frontend.png",
                "Интерфейс MITS: сократический диалог, панель рассуждений и трекинг освоения навыков",
            ),
        ],
    ),
    (
        "блок-схема общей архитектуры",
        [
            ("figures/diag_1.png", "Общая архитектура системы MITS (Frontend → Backend → AI Core)"),
        ],
    ),
    (
        "конвейера агентов",
        [
            (
                "figures/diag_2.png",
                "Мультиагентный конвейер: Profiler → Planner → Tutor → Verifier",
            ),
        ],
    ),
    (
        "sequence diagram стриминга",
        [
            ("figures/diag_7.png", "Диаграмма последовательности WebSocket-стриминга"),
        ],
    ),
    (
        "4-стадийного пайплайна",
        [
            ("figures/diag_3.png", "4-стадийный пайплайн обучения: GSPO → KTO → DPO → V-STaR-DPO"),
        ],
    ),
    (
        "базовая точность",
        [
            ("figures/fig_baseline_domains.png", "Базовая точность Qwen3.5-9B по STEM-доменам"),
        ],
    ),
    (
        "convergence plots",
        [
            (
                "figures/fig_per_stage.png",
                "Точность по стадиям обучения и доменам (Base → GSPO → KTO → DPO)",
            ),
        ],
    ),
    (
        "кривые обучения GSPO",
        [
            ("figures/01_loss_reward.png", "Динамика функции потерь и суммарной награды (GSPO)"),
            (
                "figures/02_reward_decomposition.png",
                "Декомпозиция тройной награды: корректность, формат, сократичность",
            ),
            ("figures/03_learning_rate.png", "Расписание скорости обучения (learning rate)"),
            ("figures/04_gradient_norm.png", "Норма градиента в процессе обучения GSPO"),
        ],
    ),
    (
        "гистограмма n_correct",
        [
            (
                "figures/fig_ncorrect_hist.png",
                "Бимодальное распределение n_correct (V-STaR, 426 hard-задач)",
            ),
            ("figures/fig_vstar_yield.png", "Выход within-task пар по доменам (V-STaR)"),
        ],
    ),
]


def normalize_orientation(doc) -> None:
    """Принудительно книжная (portrait) ориентация всех секций."""
    for s in doc.sections:
        if s.page_width > s.page_height:
            s.page_width, s.page_height = s.page_height, s.page_width
        s.orientation = WD_ORIENT.PORTRAIT


def place_figures(doc) -> int:
    """Заменяет маркеры [Рис ...] сопоставленными изображениями + подписями."""
    inserted = 0
    fig_no = 1
    for p in list(doc.paragraphs):
        txt = p.text.strip()
        if not (txt.startswith("[Рис") and txt.endswith("]")):
            continue
        for key, imgs in FIGURES:
            if key in txt:
                for img_rel, caption in imgs:
                    full = ROOT / img_rel
                    if not full.exists():
                        continue
                    pic_p = p.insert_paragraph_before()
                    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pic_p.add_run().add_picture(str(full), width=Inches(6.0))
                    cap_p = p.insert_paragraph_before(f"Рисунок {fig_no} — {caption}")
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in cap_p.runs:
                        r.italic = True
                    fig_no += 1
                    inserted += 1
                p._element.getparent().remove(p._element)
                break
    return inserted


def assemble() -> None:
    doc = build_document()
    pandoc = find_pandoc()
    populated = []
    for heading, sid in SECTION_IDS.items():
        md = CHAPTERS / f"{sid}.md"
        if md.exists() and md.read_text(encoding="utf-8").strip():
            frag = md_to_fragment(md, pandoc)
            n = merge_fragment(doc, heading, Document(str(frag)))
            populated.append(f"{sid}(+{n})")
    normalize_orientation(doc)
    figs = place_figures(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"[ok] assembled -> {OUT}")
    print(f"  наполнено секций: {len(populated)}; вставлено рисунков: {figs}")


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    assemble()

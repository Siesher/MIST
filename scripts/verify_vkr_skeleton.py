"""Verify ВКР skeleton: re-open .docx и проверить структуру/стили/счётчики."""

from __future__ import annotations

import io
import sys
from pathlib import Path

from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

OUT = Path(__file__).parent.parent / "docs/diploma/ВКР_Сухацкий_2026.docx"


def main() -> int:
    d = Document(str(OUT))
    txt = [p.text.strip() for p in d.paragraphs]
    joined = "\n".join(txt)
    h1 = [p.text.strip() for p in d.paragraphs if p.style and p.style.name == "Heading 1"]
    h2 = [p.text.strip() for p in d.paragraphs if p.style and p.style.name == "Heading 2"]
    body = d.element.body
    kids = [c.tag.split("}")[-1] for c in body.iterchildren()]

    checks = {
        "тема на титульнике": "мультиагентной архитектуры" in joined,
        "студент": "Сухацкий М. О." in joined,
        "руководитель": "Корлякова М. О." in joined,
        "год 2026 (не 2025)": "2026" in joined and "2025 год" not in joined,
        "АННОТАЦИЯ": "АННОТАЦИЯ" in h1,
        "СОДЕРЖАНИЕ": "СОДЕРЖАНИЕ" in h1,
        "ВВЕДЕНИЕ": "ВВЕДЕНИЕ" in h1,
        "4 главы": sum(1 for t in h1 if t.startswith("Глава")) == 4,
        "ЗАКЛЮЧЕНИЕ": "ЗАКЛЮЧЕНИЕ" in h1,
        "СПИСОК ЛИТЕРАТУРЫ": "СПИСОК ЛИТЕРАТУРЫ" in h1,
        "5 приложений": sum(1 for t in h1 if t.startswith("ПРИЛОЖЕНИЕ")) == 5,
        "H2 >= 32": len(h2) >= 32,
        "контент перед sectPr": kids[-1] == "sectPr",
    }
    ok = all(checks.values())
    for name, val in checks.items():
        print(f"  [{'OK' if val else 'FAIL'}] {name}")
    print(f"\nH1={len(h1)} H2={len(h2)} paras={len(txt)}")
    print("RESULT:", "PASS" if ok else "FAIL")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
